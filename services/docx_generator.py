from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services.schemas import ReferatContent, ReferatMeta, chapter_label, labels

FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(14)
HEADING_SIZE = Pt(16)
BLACK = RGBColor(0, 0, 0)
FIRST_LINE_INDENT = Cm(1.25)


def _apply_font(rpr_owner, name: str = FONT_NAME) -> None:
    """Force a font for all scripts and drop theme font overrides (Calibri/Cambria)."""
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        key = qn(attr)
        if key in rfonts.attrib:
            del rfonts.attrib[key]
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            key = qn(attr)
            if key in color.attrib:
                del color.attrib[key]


def _add_field(paragraph, instruction: str) -> None:
    def fld_char(kind: str):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run = paragraph.add_run()
    run._r.append(fld_char("begin"))
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    run._r.append(fld_char("separate"))
    paragraph.add_run("1")
    run = paragraph.add_run()
    run._r.append(fld_char("end"))
    for r in paragraph.runs:
        r.font.size = Pt(12)
        _apply_font(r._r)


class ReferatBuilder:
    def __init__(self, content: ReferatContent, meta: ReferatMeta, watermark: str | None = None) -> None:
        self.content = content
        self.meta = meta
        self.watermark = watermark
        self.l = labels(meta.language)
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    # ------------------------------------------------------------------ setup

    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)
        section.different_first_page_header_footer = True

        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(footer_p, "PAGE")

        first_p = section.first_page_footer.paragraphs[0]
        first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = first_p.add_run(f"{self.meta.city} – {self.meta.year}")
        run.bold = True
        run.font.size = BODY_SIZE
        _apply_font(run._r)

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = FONT_NAME
        normal.font.size = BODY_SIZE
        normal.font.color.rgb = BLACK
        _apply_font(normal.element)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

        for name, align in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT)):
            style = styles[name]
            style.font.name = FONT_NAME
            style.font.size = HEADING_SIZE
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = BLACK
            _apply_font(style.element)
            spf = style.paragraph_format
            spf.alignment = align
            spf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            spf.space_before = Pt(12 if name == "Heading 2" else 0)
            spf.space_after = Pt(12)
            spf.first_line_indent = FIRST_LINE_INDENT if name == "Heading 2" else Cm(0)
            spf.keep_with_next = True

    # ---------------------------------------------------------------- helpers

    def _para(
        self,
        text: str = "",
        *,
        size: Pt = BODY_SIZE,
        bold: bool = False,
        italic: bool = False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before: float = 0,
        space_after: float = 0,
        single: bool = True,
        left_indent: Cm | None = None,
        first_line: Cm | None = None,
    ):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = align
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE if single else WD_LINE_SPACING.ONE_POINT_FIVE
        pf.first_line_indent = first_line if first_line is not None else Cm(0)
        if left_indent is not None:
            pf.left_indent = left_indent
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            run.font.size = size
            _apply_font(run._r)
        return p

    def _body(self, text: str) -> None:
        p = self.doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = FIRST_LINE_INDENT
        run = p.add_run(text)
        _apply_font(run._r)

    def _heading(self, text: str, level: int) -> None:
        p = self.doc.add_heading(level=level)
        run = p.add_run(text)
        run.font.color.rgb = BLACK
        run.font.size = HEADING_SIZE
        _apply_font(run._r)

    def _page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- content

    def _title_page(self) -> None:
        m = self.meta
        ministry = self.l["ministry"].get(m.institution_type, self.l["ministry"]["otm"])
        self._para(ministry, bold=True, space_after=6)
        if m.institution:
            self._para(m.institution.upper(), bold=True, space_after=0)

        topic = self.content.title
        topic_size = Pt(20) if len(topic) <= 60 else Pt(18) if len(topic) <= 110 else Pt(16)
        self._para(self.l["work"].get(m.work_type, self.l["work"]["referat"]), size=Pt(28), bold=True, space_before=150, space_after=18)
        self._para(self.l["topic"], size=Pt(16), bold=True, italic=True, space_after=4)
        self._para(topic.upper(), size=topic_size, bold=True, space_after=0)

        indent = Cm(9.0)
        blank = "_" * 14
        self._para(
            f"{self.l['done_by']} {m.student or blank}",
            align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=indent, space_before=140, space_after=8,
        )
        self._para(
            f"{self.l['checked_by']} {m.teacher or blank}",
            align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=indent,
        )
        self._page_break()

    def _contents_page(self) -> None:
        self._heading(self.l["contents"], 1)
        items: list[tuple[str, bool, int]] = [(self.l["intro"].capitalize(), True, 0)]
        for ci, chapter in enumerate(self.content.chapters):
            items.append((f"{chapter_label(self.meta.language, ci)} {chapter.title}", True, 0))
            for si, section in enumerate(chapter.sections):
                items.append((f"{ci + 1}.{si + 1}. {section.title}", False, 1))
        items.append((self.l["conclusion"].capitalize(), True, 0))
        items.append((self.l["references"].capitalize(), True, 0))

        for text, bold, level in items:
            p = self._para(
                text, bold=bold, align=WD_ALIGN_PARAGRAPH.LEFT, single=False,
                left_indent=Cm(1.0 * level),
            )
            p.paragraph_format.space_after = Pt(2)
        self._page_break()

    def _body_pages(self) -> None:
        c = self.content
        self._heading(self.l["intro"], 1)
        for text in c.introduction:
            self._body(text)

        for ci, chapter in enumerate(c.chapters):
            self._page_break()
            self._heading(f"{chapter_label(self.meta.language, ci)} {chapter.title.upper()}", 1)
            for si, section in enumerate(chapter.sections):
                self._heading(f"{ci + 1}.{si + 1}. {section.title}", 2)
                for text in section.paragraphs:
                    self._body(text)

        self._page_break()
        self._heading(self.l["conclusion"], 1)
        for text in c.conclusion:
            self._body(text)

        if c.references:
            self._page_break()
            self._heading(self.l["references"], 1)
            for i, ref in enumerate(c.references, start=1):
                p = self.doc.add_paragraph(style="Normal")
                pf = p.paragraph_format
                pf.left_indent = Cm(0.75)
                pf.first_line_indent = Cm(-0.75)
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = p.add_run(f"{i}. {ref}")
                _apply_font(run._r)

    def _watermark(self) -> None:
        p = self._para(space_before=36)
        pf = p.paragraph_format
        border = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "8"), ("w:color", "999999")):
            top.set(qn(key), value)
        border.append(top)
        p._p.get_or_add_pPr().insert_element_before(
            border,
            "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
            "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
            "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
            "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange",
        )
        pf.keep_together = True
        run = p.add_run(self.watermark)
        run.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        _apply_font(run._r)

    def build(self, output_path: Path) -> Path:
        self._title_page()
        self._contents_page()
        self._body_pages()
        if self.watermark:
            self._watermark()
        props = self.doc.core_properties
        props.title = self.content.title
        props.author = self.meta.student or "SlideCraft AI"
        props.subject = self.l["work"].get(self.meta.work_type, "")
        self.doc.save(str(output_path))
        return output_path


def build_referat(
    content: ReferatContent, meta: ReferatMeta, output_path: Path, watermark: str | None = None
) -> Path:
    """Blocking: run through asyncio.to_thread()."""
    return ReferatBuilder(content, meta, watermark).build(output_path)
