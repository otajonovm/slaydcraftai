import datetime as dt
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from services.schemas import PresentationContent, Slide, labels

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN_X = 0.8
CONTENT_W = SLIDE_W - 2 * MARGIN_X
CONTENT_BOTTOM = 6.7


@dataclass(frozen=True)
class Theme:
    key: str
    name: str
    bg: str
    surface: str
    border: str
    title: str
    text: str
    muted: str
    accent: str
    accent2: str
    on_accent: str
    dark: str
    on_dark: str
    on_dark_muted: str
    heading_font: str
    body_font: str


THEMES: dict[str, Theme] = {
    "minimal_white": Theme(
        key="minimal_white",
        name="Minimal White",
        bg="FFFFFF",
        surface="F4F5F7",
        border="E5E7EB",
        title="111827",
        text="374151",
        muted="6B7280",
        accent="0F766E",
        accent2="F59E0B",
        on_accent="FFFFFF",
        dark="111827",
        on_dark="FFFFFF",
        on_dark_muted="CBD5E1",
        heading_font="Arial",
        body_font="Arial",
    ),
    "modern_blue": Theme(
        key="modern_blue",
        name="Modern Blue",
        bg="F8FAFC",
        surface="FFFFFF",
        border="DBE4F0",
        title="0F172A",
        text="334155",
        muted="64748B",
        accent="2563EB",
        accent2="06B6D4",
        on_accent="FFFFFF",
        dark="0B2A6F",
        on_dark="FFFFFF",
        on_dark_muted="BFD3F5",
        heading_font="Segoe UI",
        body_font="Segoe UI",
    ),
    "dark_tech": Theme(
        key="dark_tech",
        name="Dark Tech",
        bg="0B1120",
        surface="131C2E",
        border="22304A",
        title="F8FAFC",
        text="CBD5E1",
        muted="8B9BB4",
        accent="22D3EE",
        accent2="A78BFA",
        on_accent="0B1120",
        dark="050A16",
        on_dark="F8FAFC",
        on_dark_muted="94A3B8",
        heading_font="Trebuchet MS",
        body_font="Segoe UI",
    ),
}
DEFAULT_THEME = "modern_blue"


def get_theme(key: str) -> Theme:
    return THEMES.get(key, THEMES[DEFAULT_THEME])


def size_for(text: str, steps: Sequence[tuple[int, int]], default: int) -> int:
    length = len(text)
    for limit, size in steps:
        if length <= limit:
            return size
    return default


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _set_alpha(shape, alpha: float) -> None:
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None or len(solid) == 0:
        return
    color = solid[0]
    for old in color.findall(qn("a:alpha")):
        color.remove(old)
    etree.SubElement(color, qn("a:alpha")).set("val", str(int(alpha * 100000)))


class _Canvas:
    def __init__(self, slide, theme: Theme) -> None:
        self.slide = slide
        self.t = theme

    def background(self, color: str) -> None:
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color)

    def box(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        color: str,
        shape=MSO_SHAPE.RECTANGLE,
        *,
        line: str | None = None,
        radius: float | None = None,
        alpha: float | None = None,
    ):
        shp = self.slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(color)
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius
        if alpha is not None:
            _set_alpha(shp, alpha)
        return shp

    def label(self, shape, text: str, size: float, color: str, font: str, bold: bool = True) -> None:
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = _rgb(color)

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        text: str | Sequence[str],
        *,
        size: float,
        color: str,
        font: str,
        bold: bool = False,
        italic: bool = False,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,
        line_spacing: float = 1.1,
        space_after: float = 0,
    ):
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        items = [text] if isinstance(text, str) else list(text)
        for index, item in enumerate(items):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            p.space_after = Pt(space_after)
            run = p.add_run()
            run.text = item
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = font
            run.font.color.rgb = _rgb(color)
        return tb

    def bullets(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        items: Sequence[str],
        *,
        size: float,
        color: str,
        bullet_color: str,
        font: str,
        bullet: str = "●",
        space_after: float = 12,
        anchor=MSO_ANCHOR.TOP,
    ):
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for index, item in enumerate(items):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.line_spacing = 1.1
            p.space_after = Pt(space_after)
            mark = p.add_run()
            mark.text = f"{bullet}  "
            mark.font.size = Pt(size * 0.8)
            mark.font.color.rgb = _rgb(bullet_color)
            mark.font.name = font
            mark.font.bold = True
            run = p.add_run()
            run.text = item
            run.font.size = Pt(size)
            run.font.color.rgb = _rgb(color)
            run.font.name = font
        return tb


class PresentationBuilder:
    def __init__(self, content: PresentationContent, theme_key: str, watermark: str | None = None) -> None:
        self.content = content
        self.watermark = watermark
        self.t = get_theme(theme_key)
        self.labels = labels(content.language)
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self._blank = self.prs.slide_layouts[6]
        self.total = len(content.slides)

    def build(self, output_path: Path) -> Path:
        renderers = {
            "title": self._title_slide,
            "bullet_points": self._bullets_slide,
            "comparison": self._comparison_slide,
            "quote": self._quote_slide,
            "summary": self._summary_slide,
        }
        for index, data in enumerate(self.content.slides, start=1):
            slide = self.prs.slides.add_slide(self._blank)
            canvas = _Canvas(slide, self.t)
            renderers.get(data.layout_type, self._bullets_slide)(canvas, data, index)
            if data.speaker_notes:
                slide.notes_slide.notes_text_frame.text = data.speaker_notes

        props = self.prs.core_properties
        props.title = self.content.presentation_title
        props.author = "SlideCraft AI"
        props.subject = self.content.presentation_title
        self.prs.save(str(output_path))
        return output_path

    # ---------------------------------------------------------------- helpers

    def _content_frame(self, cv: _Canvas, data: Slide, index: int) -> float:
        t = self.t
        cv.background(t.bg)
        cv.box(0, 0, 0.16, SLIDE_H, t.accent)
        cv.box(SLIDE_W - 1.9, -0.9, 2.6, 2.6, t.accent2, MSO_SHAPE.OVAL, alpha=0.14)
        cv.box(SLIDE_W - 1.05, 0.55, 0.32, 0.32, t.accent, MSO_SHAPE.OVAL, alpha=0.6)

        cv.text(
            MARGIN_X, 0.42, CONTENT_W - 1.4, 0.95, data.title,
            size=size_for(data.title, [(40, 32), (60, 28), (85, 24)], 22),
            color=t.title, font=t.heading_font, bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0,
        )
        cv.box(MARGIN_X, 1.42, 1.1, 0.07, t.accent)
        top = 1.8
        if data.subtitle:
            cv.text(
                MARGIN_X, 1.6, CONTENT_W, 0.45, data.subtitle,
                size=16, color=t.muted, font=t.body_font, anchor=MSO_ANCHOR.MIDDLE,
            )
            top = 2.25
        self._footer(cv, index, t.muted)
        return top

    def _footer(self, cv: _Canvas, index: int, color: str) -> None:
        t = self.t
        if self.watermark and index == self.total:
            pill = cv.box(MARGIN_X, 6.9, 6.6, 0.44, t.accent, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            cv.label(pill, f"⚡ {self.watermark}", 12, t.on_accent, t.body_font)
        else:
            cv.text(
                MARGIN_X, 6.98, 8.5, 0.3, self.content.presentation_title,
                size=10, color=color, font=t.body_font, anchor=MSO_ANCHOR.MIDDLE,
            )
        cv.text(
            SLIDE_W - MARGIN_X - 1.6, 6.98, 1.6, 0.3, f"{index:02d} / {self.total:02d}",
            size=10, color=color, font=t.body_font, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
        )

    def _decor_dark(self, cv: _Canvas) -> None:
        t = self.t
        cv.background(t.dark)
        cv.box(8.7, -1.9, 6.8, 6.8, t.accent, MSO_SHAPE.OVAL, alpha=0.22)
        cv.box(10.7, 3.9, 4.3, 4.3, t.accent2, MSO_SHAPE.OVAL, alpha=0.25)
        cv.box(9.35, 5.25, 0.36, 0.36, t.accent2, MSO_SHAPE.OVAL)
        cv.box(0, 0, SLIDE_W, 0.12, t.accent)

    # ---------------------------------------------------------------- layouts

    def _title_slide(self, cv: _Canvas, data: Slide, index: int) -> None:
        t = self.t
        self._decor_dark(cv)
        chip_text = self.labels["presentation"]
        chip_w = max(2.0, 0.16 * len(chip_text) + 0.9)
        chip = cv.box(MARGIN_X, 1.25, chip_w, 0.48, t.accent, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        cv.label(chip, chip_text, 12, t.on_accent, t.heading_font)

        title = data.title or self.content.presentation_title
        cv.text(
            MARGIN_X, 1.95, 8.1, 2.7, title,
            size=size_for(title, [(28, 50), (50, 42), (80, 36), (110, 30)], 26),
            color=t.on_dark, font=t.heading_font, bold=True, anchor=MSO_ANCHOR.BOTTOM, line_spacing=1.0,
        )
        cv.box(MARGIN_X, 4.85, 1.5, 0.08, t.accent)
        subtitle = data.subtitle or (data.bullet_points[0] if data.bullet_points else "")
        if subtitle:
            cv.text(
                MARGIN_X, 5.1, 8.1, 1.2, subtitle,
                size=size_for(subtitle, [(80, 20), (140, 17)], 15),
                color=t.on_dark_muted, font=t.body_font,
            )
        cv.text(
            MARGIN_X, 6.75, 4, 0.35, str(dt.date.today().year),
            size=12, color=t.on_dark_muted, font=t.body_font, bold=True,
        )

    def _bullets_slide(self, cv: _Canvas, data: Slide, index: int) -> None:
        t = self.t
        top = self._content_frame(cv, data, index)
        items = data.bullet_points[:6] or [data.subtitle or data.title]
        n = len(items)
        cols, rows = (n, 1) if n <= 3 else (2, 2) if n == 4 else (3, 2)
        gap = 0.3
        area_h = CONTENT_BOTTOM - top
        card_w = (CONTENT_W - gap * (cols - 1)) / cols
        card_h = (area_h - gap * (rows - 1)) / rows
        longest = max(items, key=len)

        for i, item in enumerate(items):
            r, c = divmod(i, cols)
            x = MARGIN_X + c * (card_w + gap)
            y = top + r * (card_h + gap)
            cv.box(x, y, card_w, card_h, t.surface, MSO_SHAPE.ROUNDED_RECTANGLE, line=t.border, radius=0.07)
            if rows == 1:
                cv.box(x, y + 0.35, 0.07, 0.75, t.accent)
                badge = cv.box(x + 0.4, y + 0.35, 0.75, 0.75, t.accent, MSO_SHAPE.OVAL)
                cv.label(badge, f"{i + 1:02d}", 16, t.on_accent, t.heading_font)
                base = 22 if cols == 1 else 20 if cols == 2 else 18
                size = size_for(longest, [(60, base), (100, base - 2), (150, base - 4)], base - 6)
                cv.text(
                    x + 0.4, y + 1.4, card_w - 0.8, card_h - 1.7, item,
                    size=size, color=t.text, font=t.body_font, line_spacing=1.15,
                )
            else:
                badge = cv.box(x + 0.3, y + card_h / 2 - 0.31, 0.62, 0.62, t.accent, MSO_SHAPE.OVAL)
                cv.label(badge, str(i + 1), 15, t.on_accent, t.heading_font)
                base = 18 if cols == 2 else 16
                size = size_for(longest, [(60, base), (100, base - 1), (150, base - 3)], base - 4)
                cv.text(
                    x + 1.15, y + 0.25, card_w - 1.4, card_h - 0.5, item,
                    size=size, color=t.text, font=t.body_font, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1,
                )

    def _comparison_slide(self, cv: _Canvas, data: Slide, index: int) -> None:
        t = self.t
        top = self._content_frame(cv, data, index)
        gap = 0.6
        col_w = (CONTENT_W - gap) / 2
        header_h = 0.75
        columns = [
            (data.left_title, data.left_points, t.accent),
            (data.right_title, data.right_points, t.accent2),
        ]
        all_points = data.left_points + data.right_points
        longest = max(all_points, key=len) if all_points else ""
        count = max(len(data.left_points), len(data.right_points), 1)
        size = size_for(longest, [(50, 18), (90, 16), (140, 15)], 14)
        if count >= 5:
            size -= 1

        for c, (heading, points, color) in enumerate(columns):
            x = MARGIN_X + c * (col_w + gap)
            cv.box(x, top, col_w, CONTENT_BOTTOM - top, t.surface, MSO_SHAPE.ROUNDED_RECTANGLE, line=t.border, radius=0.05)
            head = cv.box(x, top, col_w, header_h, color, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
            cv.label(
                head, heading or ("A" if c == 0 else "B"),
                size_for(heading, [(28, 20), (45, 17)], 15), t.on_accent, t.heading_font,
            )
            cv.bullets(
                x + 0.4, top + header_h + 0.35, col_w - 0.8, CONTENT_BOTTOM - top - header_h - 0.6, points,
                size=size, color=t.text, bullet_color=color, font=t.body_font, bullet="■", space_after=10,
            )

        vs = cv.box(SLIDE_W / 2 - 0.42, top + header_h / 2 - 0.42 + 0.05, 0.84, 0.84, t.title, MSO_SHAPE.OVAL, line=t.bg)
        cv.label(vs, "VS", 14, t.bg, t.heading_font)

    def _quote_slide(self, cv: _Canvas, data: Slide, index: int) -> None:
        t = self.t
        cv.background(t.bg)
        cv.box(0, 0, 0.16, SLIDE_H, t.accent)
        cv.box(-1.2, 5.2, 3.2, 3.2, t.accent2, MSO_SHAPE.OVAL, alpha=0.15)
        cv.box(SLIDE_W - 2.0, -1.0, 3.0, 3.0, t.accent, MSO_SHAPE.OVAL, alpha=0.12)
        cv.text(
            MARGIN_X, 0.45, CONTENT_W, 0.7, data.title,
            size=22, color=t.muted, font=t.heading_font, bold=True, anchor=MSO_ANCHOR.MIDDLE,
        )

        card_x, card_y, card_w, card_h = 1.3, 1.45, SLIDE_W - 2.6, 4.85
        cv.box(card_x, card_y, card_w, card_h, t.surface, MSO_SHAPE.ROUNDED_RECTANGLE, line=t.border, radius=0.06)
        cv.box(card_x, card_y + 0.6, 0.1, card_h - 1.2, t.accent)
        cv.text(
            card_x + 0.5, card_y - 0.35, 1.6, 1.6, "\u201C",
            size=130, color=t.accent, font="Georgia", bold=True,
        )
        quote = data.quote or data.subtitle
        cv.text(
            card_x + 1.0, card_y + 0.85, card_w - 2.0, card_h - 2.0, quote,
            size=size_for(quote, [(80, 30), (140, 26), (220, 22)], 19),
            color=t.title, font=t.heading_font, italic=True, anchor=MSO_ANCHOR.MIDDLE,
            align=PP_ALIGN.CENTER, line_spacing=1.15,
        )
        if data.quote_author:
            cv.text(
                card_x + 1.0, card_y + card_h - 1.0, card_w - 2.0, 0.5, f"— {data.quote_author}",
                size=18, color=t.accent, font=t.body_font, bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        self._footer(cv, index, t.muted)

    def _summary_slide(self, cv: _Canvas, data: Slide, index: int) -> None:
        t = self.t
        self._decor_dark(cv)
        title = data.title or self.labels["summary"]
        cv.text(
            MARGIN_X, 0.5, 7.6, 1.0, title,
            size=size_for(title, [(30, 36), (55, 30)], 26),
            color=t.on_dark, font=t.heading_font, bold=True, anchor=MSO_ANCHOR.MIDDLE,
        )
        cv.box(MARGIN_X, 1.55, 1.3, 0.08, t.accent)

        items = (data.bullet_points or [data.subtitle or title])[:5]
        area_top, area_h = 1.95, 4.65
        row_gap = 0.18
        row_h = min(1.15, (area_h - row_gap * (len(items) - 1)) / len(items))
        longest = max(items, key=len)
        size = size_for(longest, [(60, 18), (100, 16), (150, 14)], 13)
        for i, item in enumerate(items):
            y = area_top + i * (row_h + row_gap)
            cv.box(MARGIN_X, y, 7.5, row_h, t.on_dark, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18, alpha=0.08)
            check = cv.box(MARGIN_X + 0.25, y + row_h / 2 - 0.25, 0.5, 0.5, t.accent, MSO_SHAPE.OVAL)
            cv.label(check, "✓", 14, t.on_accent, t.body_font)
            cv.text(
                MARGIN_X + 1.0, y + 0.08, 6.3, row_h - 0.16, item,
                size=size, color=t.on_dark, font=t.body_font, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05,
            )

        panel = cv.box(8.95, 1.95, 3.6, 4.65, t.accent, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        panel.shadow.inherit = False
        cv.text(
            9.25, 2.35, 3.0, 2.4, self.labels["thanks"],
            size=26, color=t.on_accent, font=t.heading_font, bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05,
        )
        cv.box(10.35, 4.85, 0.8, 0.06, t.on_accent)
        cv.text(
            9.25, 5.05, 3.0, 1.1, self.labels["questions"],
            size=18, color=t.on_accent, font=t.body_font, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        self._footer(cv, index, t.on_dark_muted)


def build_presentation(
    content: PresentationContent, theme_key: str, output_path: Path, watermark: str | None = None
) -> Path:
    """Blocking: run through asyncio.to_thread()."""
    return PresentationBuilder(content, theme_key, watermark).build(output_path)
